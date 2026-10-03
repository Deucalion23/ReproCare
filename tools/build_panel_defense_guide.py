from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\xampp\htdocs\ReproCare\docs\REPROCARE_PANEL_DEFENSE_GUIDE.docx"

NAVY = "17365D"
BLUE = "2E74B5"
INK = "1F2937"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
GOLD = "A66B00"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_run_font(run, size=11, bold=None, color=INK, italic=False, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_text(doc, text="", style="Normal", before=0, after=6, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p, r


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, 10.5)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    lead = p.add_run(label + " ")
    set_run_font(lead, 10.5, True, NAVY)
    body = p.add_run(text)
    set_run_font(body, 10.5, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers, rows, widths, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, font_size, True, NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, row_data)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            r = p.add_run(text)
            set_run_font(r, font_size, index == 0, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("ReproCare Panel Defense Guide | ")
    set_run_font(r, 8, False, MUTED)
    r = p.add_run("Page ")
    set_run_font(r, 8, False, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Cover page - editorial cover pattern, compact reference guide styling.
    for _ in range(7):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    p, r = add_text(doc, "PANEL DEFENSE GUIDE", before=0, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(r, 12, True, GOLD)
    p, r = add_text(doc, "ReproCare", before=0, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(r, 30, True, NAVY)
    p, r = add_text(doc, "Maternal and Reproductive Health Tracking and Learning System for Community Well-Being", before=0, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(r, 15, False, BLUE)
    p, r = add_text(doc, "Presentation explanation script, evidence guide, and possible panel questions", before=0, after=40, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(r, 11, False, MUTED, italic=True)
    p, r = add_text(doc, "Prepared for the Pre-Final Defense", before=0, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(r, 11, True, INK)
    p, r = add_text(doc, "Pangasinan State University - San Carlos Campus", before=0, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(r, 10.5, False, MUTED)
    p, r = add_text(doc, "October 2026", before=0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(r, 10.5, False, MUTED)
    add_page_break(doc)

    add_heading(doc, "How to Use This Guide", 1)
    p, r = add_text(doc, "This guide is a speaking aid, not a script that must be memorized word for word. Use the short answer first, then add evidence or a limitation only when the panel asks for more detail.")
    set_run_font(r)
    add_callout(doc, "Defense pattern:", "Answer directly -> explain the evidence -> state the boundary or limitation -> return to the project's value.")
    add_heading(doc, "Core Defense Message", 2)
    p, r = add_text(doc, "ReproCare digitizes a community maternal and reproductive health workflow that was previously dependent on paper records, manual follow-up, and fragmented reporting. It gives each authorized role a focused view of the same coordinated process: data entry, clinical review, follow-up, reporting, and health education.")
    set_run_font(r)
    add_heading(doc, "Suggested Opening Statement (about 60-90 seconds)", 2)
    opening = (
        "Good day, panelists. Our project is ReproCare: Maternal and Reproductive Health Tracking and Learning System for Community Well-Being. "
        "It was developed to address the difficulty of relying on paper records, delayed updates, fragmented documentation, and inconsistent follow-up in community maternal-health services. "
        "ReproCare centralizes patient information, menstrual and pregnancy tracking, checkup monitoring, health records, reports, learning materials, and role-based coordination among women, BHWs, BHW Presidents, midwives, RHU administrators, and CHO administrators. "
        "The system does not replace clinical judgment. Its purpose is to organize information, surface follow-up needs, and support licensed health personnel in making informed decisions. "
        "Based on our ISO/IEC 25010:2011 acceptability evaluation with 120 respondents, the system obtained an overall average weighted mean of 3.76, interpreted as Agree or Acceptable."
    )
    p, r = add_text(doc, opening, after=8)
    set_run_font(r)
    add_heading(doc, "Key Numbers to Remember", 2)
    add_matrix(doc, ["Item", "Result to State"], [
        ("Respondents", "120 total respondents"),
        ("Women users", "100 respondents (83.33%)"),
        ("BHWs", "16 respondents (13.33%)"),
        ("Other roles", "One CHO/Super Admin, one RHU Admin, one Midwife, and one BHW President"),
        ("Overall evaluation", "3.76 Average Weighted Mean; Agree / Acceptable"),
        ("Highest criterion", "Security: 3.83"),
        ("Usability", "3.81"),
        ("Functional suitability", "3.80"),
    ], [2700, 6660], 10)
    add_page_break(doc)

    add_heading(doc, "Slide-by-Slide Explanation Guide", 1)
    slides = [
        ("1. Title", "Introduce the system and the team.", "Our project is ReproCare, a web-based and mobile-responsive maternal and reproductive health tracking and learning system designed for community-based healthcare coordination."),
        ("2. Objectives", "Connect every objective to a feature.", "We analyzed the manual workflow, designed a centralized secure system, integrated smart decision-support capability, and evaluated its acceptability among intended users."),
        ("3. Methodology", "Explain why Agile Scrum was appropriate.", "Agile Scrum allowed us to develop the project incrementally, review modules in short cycles, collect feedback, and adjust functions based on stakeholder needs."),
        ("4. User Requirements", "Emphasize minimum necessary access.", "The system uses role-based access because each user has a different responsibility. Women monitor their own information; BHWs encode field data; supervisors review; and administrators manage reports, users, and oversight."),
        ("5. Technology Used", "Explain the stack by layer.", "Laravel and PHP manage the application logic, SQLite is the current local data store, and Bootstrap, Tailwind CSS, JavaScript, Chart.js, and Leaflet support the user interface, charts, and maps."),
        ("6-7. Manual Process and Challenges", "Describe the problem without overstating it.", "The previous workflow relied heavily on paper forms and manual compilation. This could delay updates, make retrieval difficult, fragment records, and make coordinated follow-up more difficult."),
        ("8. Proposed Process", "Show the end-to-end flow.", "The proposed process connects data entry, review, follow-up, reporting, education, and communication in one role-based platform. The system is intended to improve organization and continuity of care."),
        ("9. AI Integration", "State the safeguard before describing the benefit.", "The smart-suggestion capability is supplementary decision support. It uses grouped analytics and system rules or an optional AI provider to summarize priorities; it does not diagnose, prescribe, or replace a licensed clinician."),
        ("10-12. Evaluation and Testing", "Lead with sample and result.", "We evaluated acceptability using ISO/IEC 25010:2011 with 120 respondents. The overall AWM was 3.76, interpreted as Agree or Acceptable, showing positive user assessment across the evaluated quality criteria."),
        ("13. Closing", "End with the contribution.", "ReproCare contributes a structured, role-based way to support maternal-health tracking, patient education, follow-up, and reporting at the community level."),
    ]
    add_matrix(doc, ["Slide", "What to Emphasize", "Suggested Explanation"], slides, [1500, 2600, 5260], 8.8)
    add_page_break(doc)

    add_heading(doc, "System Flow You Can Explain", 1)
    add_number(doc, "A woman registers and supplies the required profile information. The account is subject to email and RHU approval controls before protected patient functions are available.")
    add_number(doc, "A BHW can register or assist patients, record community health information, and submit records or reports for review.")
    add_number(doc, "The BHW President and Midwife review records and reports within their assigned coverage and return incomplete information for correction when necessary.")
    add_number(doc, "The RHU Administrator manages approval queues, staff records, reports, maternal health administration, and selected communication functions.")
    add_number(doc, "The CHO Administrator views city-wide monitoring information, analytics, audit logs, and system-level administrative functions.")
    add_number(doc, "Patients can view their information, track cycles and pregnancies, access learning materials, and receive notifications after they meet access requirements.")
    add_callout(doc, "One-sentence explanation:", "ReproCare transforms a fragmented paper workflow into a controlled digital flow where each role records, reviews, approves, or uses the information appropriate to its responsibility.")
    add_heading(doc, "Role Explanation at a Glance", 2)
    add_matrix(doc, ["User Role", "Primary Responsibility", "Examples of Functions"], [
        ("Woman / Patient", "Self-monitoring and participation", "Cycle and pregnancy tracking, checkup viewing, learning materials, notifications, messaging"),
        ("BHW", "Community data collection and follow-up", "Patient and walk-in registration, health records, pregnancy data, report submission"),
        ("BHW President", "Barangay-level supervision", "Review of BHW records, reports, tasks, and transfer requests"),
        ("Midwife", "Clinical review and coordination", "Checkup management, referral review, risk monitoring, report review"),
        ("RHU Administrator", "Facility-level administration", "Patient approval, staff management, reports, maternal events, analytics"),
        ("CHO Administrator", "City-wide oversight", "City analytics, audits, account management, reports, backup management"),
    ], [1700, 2700, 4960], 8.8)
    add_page_break(doc)

    add_heading(doc, "Evaluation Results You Should Be Ready to Explain", 1)
    p, r = add_text(doc, "The study reports an ISO/IEC 25010:2011-based acceptability evaluation. The correct defense approach is to explain the result as respondent perception of the evaluated system, not as proof of clinical effectiveness or improved health outcomes.")
    set_run_font(r)
    results = [
        ("Functional Suitability", "3.80", "Agree / Acceptable", "Respondents assessed the availability and appropriateness of the functions."),
        ("Performance Efficiency", "3.79", "Agree / Acceptable", "Respondents assessed response and operation during use."),
        ("Compatibility", "3.72", "Agree / Acceptable", "Respondents assessed how the system worked with related functions and use contexts."),
        ("Usability", "3.81", "Agree / Acceptable", "Respondents assessed clarity, learnability, and navigation."),
        ("Reliability", "3.70", "Agree / Acceptable", "Respondents assessed consistency and dependable operation."),
        ("Security", "3.83", "Agree / Acceptable", "Respondents assessed access control, authentication, and protection-oriented functions."),
        ("Maintainability", "3.68", "Agree / Acceptable", "Respondents assessed organization and understandability of the system."),
        ("Portability", "3.71", "Agree / Acceptable", "Respondents assessed use across devices and screen sizes."),
        ("Overall", "3.76", "Agree / Acceptable", "The overall respondent evaluation was positive."),
    ]
    add_matrix(doc, ["Criterion", "AWM", "Interpretation", "How to Explain It"], results, [2100, 900, 1800, 4560], 8.4)
    add_callout(doc, "Safe interpretation:", "The findings show that respondents generally agreed that ReproCare was acceptable based on the evaluated quality criteria. The evaluation does not by itself establish that the system reduces maternal mortality, makes diagnoses, or replaces professional care.", "FFF6E8")
    add_page_break(doc)

    add_heading(doc, "Likely Panel Questions and Suggested Answers", 1)
    qa_one = [
        ("1. What problem does ReproCare solve?", "It addresses fragmented paper-based documentation, delayed updates, manual report preparation, and difficult follow-up by centralizing maternal and reproductive health workflows in one role-based system."),
        ("2. Why is the system web-based and mobile-responsive?", "Healthcare workers and women may use different devices. A web-based responsive system can be accessed through a browser on desktop, tablet, or smartphone without requiring a separate native application installation."),
        ("3. Who are the intended users?", "The intended users are women, BHWs, BHW Presidents, Midwives, RHU Administrators, and CHO/Super Administrators. Each role receives only the functions appropriate to its responsibility."),
        ("4. What makes the system different from ordinary record keeping?", "It does not only digitize records. It connects records with pregnancy and cycle tracking, checkups, review workflows, notifications, reports, learning materials, and role-based coordination."),
        ("5. How does ReproCare protect privacy?", "The system uses authenticated accounts, role-based access control, password hashing, input validation, session protection, activity logging, and approval controls. Access is limited according to assigned responsibilities."),
        ("6. Does the system make medical diagnoses?", "No. ReproCare organizes records and surfaces follow-up priorities. Any clinical assessment, diagnosis, prescription, or treatment decision remains the responsibility of qualified healthcare professionals."),
        ("7. What is the role of AI in the system?", "AI or rules-based smart suggestions are supplementary decision support for grouped analytics. They help summarize operational priorities, but they cannot change records, create appointments, prescribe treatment, or replace clinical judgment."),
        ("8. Why use Agile Scrum?", "The system has several users and workflows. Agile Scrum allowed incremental development, frequent review, and adjustments after feedback rather than waiting until the end of the project to validate all functions."),
    ]
    add_matrix(doc, ["Possible Question", "Suggested Answer"], qa_one, [3000, 6360], 9)
    add_page_break(doc)

    add_heading(doc, "Likely Panel Questions and Suggested Answers (continued)", 1)
    qa_two = [
        ("9. How was the system evaluated?", "The system was evaluated using ISO/IEC 25010:2011 quality criteria. There were 120 respondents, and the overall Average Weighted Mean was 3.76, interpreted in the study as Agree or Acceptable."),
        ("10. What does the score of 3.76 mean?", "It means respondents generally agreed that the system met the evaluated quality criteria. It is an acceptability result, not a measure of clinical outcomes or a guarantee that every future user will have the same experience."),
        ("11. Why were most respondents women users?", "Women are the largest direct user group of the system. The remaining respondents represented the healthcare and supervisory roles that use specialized administrative and clinical workflow features."),
        ("12. What are the limitations of the system?", "The system is limited to maternal and reproductive health tracking, education, follow-up, reporting, and related workflows. It is not a hospital information system, telemedicine service, emergency-response system, national database, or replacement for clinical judgment."),
        ("13. Does the system work without internet?", "Core authenticated health workflows require connectivity. The application includes limited PWA caching and selected queued synchronization support, but it should not be presented as a full offline clinical-record system."),
        ("14. Can the system integrate with hospitals or national databases?", "Not in the present scope. The current implementation is designed for the identified community workflow. External hospital and national-database integration are possible future enhancements subject to interoperability, privacy, and institutional approval requirements."),
        ("15. How are duplicate patients handled?", "The RHU approval process includes a possible-duplicate review workflow. This helps reviewers identify likely duplicate registrations before approving or linking relevant patient information."),
        ("16. What happens if a patient enters incorrect information?", "The workflow supports review, correction, return, and resubmission. Health personnel verify information as part of their assigned responsibilities rather than treating all self-entered information as clinically verified."),
    ]
    add_matrix(doc, ["Possible Question", "Suggested Answer"], qa_two, [3000, 6360], 9)
    add_page_break(doc)

    add_heading(doc, "Technical Questions and Honest Answers", 1)
    qa_tech = [
        ("What technologies were used?", "The current local implementation uses PHP and Laravel, SQLite, Blade, HTML, CSS, JavaScript, Bootstrap, Tailwind CSS, Vite, Chart.js, Leaflet, TCPDF, and PWA web technologies. The full technology list should distinguish active technologies from optional integrations."),
        ("Is the database MySQL?", "For the current local demonstration, the configured database is SQLite. Laravel also supports MySQL, MariaDB, and PostgreSQL. The presentation slide should be updated if it currently states MySQL as the active demonstration database."),
        ("Is Google Sign-In already live?", "The Google Sign-In workflow has been implemented using Laravel Socialite, but it becomes live only after Google OAuth credentials are configured in the deployment environment."),
        ("Are SMS reminders currently delivered?", "The system has TextBee and Movider SMS integration logic. Actual delivery requires a configured provider account, API key, device or gateway, and network access. During a demonstration, we should state whether those credentials are configured."),
        ("Does email verification send real emails now?", "The verification workflow is implemented. In the current local configuration, mail uses a log driver, so a production SMTP or transactional email provider must be configured for real inbox delivery."),
        ("Does the system use Qwen AI in the demo?", "The code supports optional Groq-based analytics with a configurable model, and it also has a rules-based fallback. The current local configuration uses rules-based analytics unless a Groq key and provider setting are configured. We should not claim a live AI provider unless it is configured for the defense."),
        ("Why not allow every role to use Google Sign-In?", "Patient Google sign-in is designed as a convenience onboarding option. Staff accounts remain internally provisioned to preserve tighter administrative and clinical access control."),
        ("How is data backed up?", "The system includes administrator-controlled database export, import, download, and backup management functions. In production, backup frequency and retention should follow the facility's approved policy."),
    ]
    add_matrix(doc, ["Technical Question", "Accurate Answer"], qa_tech, [3000, 6360], 8.9)
    add_callout(doc, "Important:", "Never claim that an optional external integration is live unless it is configured and successfully demonstrated. Use the terms implemented, supported, or optional when that is the accurate status.", "FDEDEC")
    add_page_break(doc)

    add_heading(doc, "Questions About Data Privacy, Safety, and Scope", 1)
    qa_scope = [
        ("Why should patients trust the system with health data?", "The system applies authentication, role-based access, password hashing, validation, approval workflows, and activity logging. However, deployment must also follow the health facility's data privacy, retention, access-control, and operational policies."),
        ("Can patients see other patients' data?", "No. Patient functions are scoped to the authenticated patient's own records. Staff access is role- and coverage-based."),
        ("Can BHWs approve their own entries?", "No. The workflow separates data entry from review. Records and reports can be reviewed by supervisory roles, and patient account approval belongs to authorized RHU administration."),
        ("What if the AI recommendation is wrong?", "The system treats smart suggestions as non-diagnostic operational support. The interface and documentation emphasize that final decisions remain with licensed health professionals."),
        ("Why include a forum in a health system?", "The forum supports community health awareness and discussion. It does not replace private consultation, diagnosis, or emergency services. Moderation and responsible-use practices remain important."),
        ("What is outside the scope?", "General hospital management, telemedicine consultations, emergency dispatch, national health-database integration, and automatic clinical diagnosis are outside the current project scope."),
    ]
    add_matrix(doc, ["Possible Question", "Suggested Answer"], qa_scope, [3000, 6360], 9)
    add_heading(doc, "Questions to Answer Carefully", 2)
    add_bullet(doc, "Do not claim that ReproCare has proven to reduce maternal mortality. State that it is designed to support organized tracking, follow-up, and decision support.")
    add_bullet(doc, "Do not call analytics or AI a diagnostic tool. State that it is supplementary and requires human review.")
    add_bullet(doc, "Do not claim fully offline clinical operation. State that core workflows require connectivity, while the application has limited PWA support.")
    add_bullet(doc, "Do not claim live SMS, live Google Sign-In, live Groq AI, or real email delivery unless the required credentials are configured and tested.")
    add_page_break(doc)

    add_heading(doc, "Items to Align Before the Defense", 1)
    p, r = add_text(doc, "The paper, slides, and current code should use consistent descriptions. These items are not flaws in the research; they are presentation details that should be aligned before the panel review.")
    set_run_font(r)
    alignment_rows = [
        ("Database", "The presentation identifies MySQL, while the current local application configuration uses SQLite.", "For the actual demonstration, state SQLite. If using MySQL in another deployment, describe it as supported or deployment-specific."),
        ("AI provider", "The slide names Qwen 3.8-27B, while the local configuration defaults to rules-based analytics unless Groq is configured.", "Demonstrate the actual active mode. Call Qwen/Groq optional unless the API key, provider setting, and successful test are present."),
        ("SMS", "The system has provider integrations, but live delivery depends on credentials and a gateway.", "Say SMS capability is implemented; demonstrate delivery only if configured."),
        ("Email verification", "The local mail driver records messages in logs rather than sending to an external inbox.", "Say the verification workflow is implemented; production delivery requires SMTP or a transactional provider."),
        ("Response-scale legend", "The supplied paper and presentation use different upper-band ranges for Strongly Agree.", "Use one scale consistently in the oral presentation, tables, and slides. Confirm the approved scale with the research adviser before printing."),
        ("Offline statement", "The paper describes no offline support, while the current code includes limited PWA caching and sync support.", "State that the system is not a full offline clinical system; the PWA provides limited resilience only."),
    ]
    add_matrix(doc, ["Item", "What to Check", "Safe Defense Position"], alignment_rows, [1500, 3350, 4510], 8.6)
    add_page_break(doc)

    add_heading(doc, "Strong Closing Statement", 1)
    closing = (
        "In conclusion, ReproCare was developed to support a more organized and coordinated maternal and reproductive health workflow at the community level. "
        "It brings together patient tracking, checkups, records, reports, learning materials, communication, and role-based accountability in one web-based system. "
        "Our evaluation results show that respondents generally found the system acceptable across the ISO/IEC 25010:2011 quality criteria. "
        "Most importantly, ReproCare is designed to support healthcare personnel and patients; it does not replace the judgment, responsibility, or care of licensed professionals. Thank you."
    )
    p, r = add_text(doc, closing, after=14)
    set_run_font(r, 11.5)
    add_heading(doc, "Final Defense Checklist", 2)
    checklist = [
        "Confirm the exact database used in the live demonstration.",
        "Confirm whether SMS, Google Sign-In, Groq AI, and real email delivery are configured; do not claim them as live if they are not.",
        "Use one approved Likert-scale legend consistently across slides and paper.",
        "Prepare one patient account and one staff account for a controlled feature demonstration.",
        "Demonstrate a clear workflow: patient or walk-in entry -> health record -> review/report -> dashboard or analytics view.",
        "Keep the AI explanation limited to supplementary decision support and human review.",
        "Know the overall AWM (3.76), sample size (120), and the highest criterion (Security, 3.83).",
        "Be ready to explain scope limitations before the panel asks.",
    ]
    for item in checklist:
        add_bullet(doc, item)
    add_callout(doc, "Best final habit:", "If you do not know an answer, do not guess. Say: 'Based on the current scope, that is not implemented yet, but it is a valid future enhancement.'", "EAF4EA")

    doc.core_properties.title = "ReproCare Panel Defense Guide"
    doc.core_properties.subject = "Explanation script and possible panel questions"
    doc.core_properties.author = "ReproCare Research Team"
    doc.save(OUT)


if __name__ == "__main__":
    build()
